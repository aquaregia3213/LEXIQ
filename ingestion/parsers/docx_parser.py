import docx
from typing import List, Dict, Any
from ingestion.preprocessing.cleaner import clean_extracted_text


class DOCXParsingError(Exception):
    """Exception raised when DOCX parsing fails."""
    pass


def parse_docx_document(file_path: str) -> Dict[str, Any]:
    """
    Parse DOCX document using python-docx.
    Returns dict containing page_count, pages list, and full_text.
    """
    try:
        document = docx.Document(file_path)
    except Exception as err:
        raise DOCXParsingError(f"python-docx failed to open DOCX document: {err}")

    try:
        paragraphs_text: List[str] = []

        # Extract text from main paragraphs and headings
        for para in document.paragraphs:
            cleaned = clean_extracted_text(para.text)
            if cleaned:
                paragraphs_text.append(cleaned)

        # Extract text from tables if present
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join([clean_extracted_text(cell.text) for cell in row.cells if cell.text.strip()])
                if row_text:
                    paragraphs_text.append(row_text)

        full_text = "\n\n".join(paragraphs_text)

        # Estimate page count (approx ~350 words or 2200 chars per page for DOCX without page break tags)
        estimated_pages = max(1, (len(full_text) // 2200) + 1)

        # Structure page chunks
        pages_data = [{
            "page_number": 1,
            "text": full_text,
            "char_count": len(full_text)
        }]

        return {
            "file_type": "DOCX",
            "page_count": estimated_pages,
            "pages": pages_data,
            "full_text": full_text
        }
    except Exception as err:
        raise DOCXParsingError(f"Error parsing DOCX content: {err}")
