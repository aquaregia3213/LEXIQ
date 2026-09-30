import os
import uuid
import tempfile
import fitz  # PyMuPDF
import docx
import pytest
from sqlalchemy import create_engine
from sqlalchemy.pool import StaticPool
from sqlalchemy.orm import sessionmaker
from fastapi.testclient import TestClient

from backend.app.database import Base, get_db
from backend.app.main import app
from ingestion.preprocessing.validator import validate_uploaded_file, FileValidationError
from ingestion.parsers.pdf_parser import parse_pdf_document
from ingestion.parsers.docx_parser import parse_docx_document
from ingestion.pipeline.ingestor import process_contract_ingestion


@pytest.fixture
def db_session():
    """Isolated in-memory SQLite session fixture using StaticPool."""
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool
    )
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(bind=engine)


@pytest.fixture
def test_pdf_file():
    """Generate temporary 2-page PDF file with structured sections."""
    pdf_filename = f"test_{uuid.uuid4().hex}.pdf"
    pdf_path = os.path.join(tempfile.gettempdir(), pdf_filename)
    
    doc = fitz.open()

    # Page 1
    page1 = doc.new_page()
    page1.insert_text((50, 50), "Non-Disclosure Agreement\n\nSection 1: Confidentiality\nThis agreement governs confidential information shared between Acme Corp and TechProvider Inc.")
    
    # Page 2
    page2 = doc.new_page()
    page2.insert_text((50, 50), "Section 2: Term and Termination\nEither party may terminate this agreement upon 30 days notice.")

    doc.save(pdf_path)
    doc.close()

    yield pdf_path

    if os.path.exists(pdf_path):
        try:
            os.remove(pdf_path)
        except Exception:
            pass


@pytest.fixture
def test_docx_file():
    """Generate temporary DOCX file with structured sections."""
    docx_filename = f"test_{uuid.uuid4().hex}.docx"
    docx_path = os.path.join(tempfile.gettempdir(), docx_filename)

    doc = docx.Document()
    doc.add_heading("Master Services Agreement", level=1)
    doc.add_paragraph("Section 1: Scope of Work. Contractor shall perform software engineering services for Client.")
    doc.add_paragraph("Section 2: Governing Law. This agreement is governed by the laws of California.")
    doc.save(docx_path)

    yield docx_path

    if os.path.exists(docx_path):
        try:
            os.remove(docx_path)
        except Exception:
            pass


def test_pdf_parsing_and_page_tracking(test_pdf_file):
    """Verify PyMuPDF parser extracts text and preserves page numbers."""
    parsed = parse_pdf_document(test_pdf_file)
    assert parsed["file_type"] == "PDF"
    assert parsed["page_count"] == 2
    assert len(parsed["pages"]) == 2
    assert parsed["pages"][0]["page_number"] == 1
    assert "Section 1: Confidentiality" in parsed["pages"][0]["text"]
    assert parsed["pages"][1]["page_number"] == 2
    assert "Section 2: Term and Termination" in parsed["pages"][1]["text"]


def test_docx_parsing(test_docx_file):
    """Verify python-docx parser extracts headings and paragraph text."""
    parsed = parse_docx_document(test_docx_file)
    assert parsed["file_type"] == "DOCX"
    assert parsed["page_count"] >= 1
    assert "Master Services Agreement" in parsed["full_text"]
    assert "Section 1: Scope of Work" in parsed["full_text"]


def test_file_validation():
    """Test file validator error handling for size, extension, and corrupted files."""
    # Test 1: Empty file
    empty_filename = f"empty_{uuid.uuid4().hex}.pdf"
    empty_path = os.path.join(tempfile.gettempdir(), empty_filename)
    with open(empty_path, "wb") as f:
        f.write(b"")

    try:
        with pytest.raises(FileValidationError, match="0 bytes"):
            validate_uploaded_file(empty_path, "empty.pdf")
    finally:
        if os.path.exists(empty_path):
            os.remove(empty_path)

    # Test 2: Invalid extension
    txt_filename = f"txt_{uuid.uuid4().hex}.txt"
    txt_path = os.path.join(tempfile.gettempdir(), txt_filename)
    with open(txt_path, "wb") as f:
        f.write(b"some plain text")

    try:
        with pytest.raises(FileValidationError, match="Unsupported file extension"):
            validate_uploaded_file(txt_path, "invalid.txt")
    finally:
        if os.path.exists(txt_path):
            os.remove(txt_path)

    # Test 3: Corrupted PDF header
    bad_pdf_filename = f"bad_{uuid.uuid4().hex}.pdf"
    bad_pdf_path = os.path.join(tempfile.gettempdir(), bad_pdf_filename)
    with open(bad_pdf_path, "wb") as f:
        f.write(b"NOT_A_REAL_PDF_HEADER")

    try:
        with pytest.raises(FileValidationError, match="Corrupted PDF file"):
            validate_uploaded_file(bad_pdf_path, "corrupt.pdf")
    finally:
        if os.path.exists(bad_pdf_path):
            os.remove(bad_pdf_path)


def test_full_ingestion_pipeline_pdf(db_session, test_pdf_file):
    """Test full ingestion pipeline processing for a PDF contract file."""
    with open(test_pdf_file, "rb") as f:
        file_bytes = f.read()

    result = process_contract_ingestion(
        db=db_session,
        file_bytes=file_bytes,
        original_filename="nda_contract.pdf",
        contract_name="NDA Agreement",
        contract_type="Non-Disclosure Agreement"
    )

    assert result["status"] == "success"
    assert result["contract_name"] == "NDA Agreement"
    assert result["file_type"] == "PDF"
    assert result["page_count"] == 2
    assert result["chunk_count"] >= 2
    assert result["processing_status"] == "COMPLETED"


def test_upload_api_endpoint(db_session, test_pdf_file):
    """Test FastAPI REST endpoint POST /api/contracts/upload."""
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)

    try:
        with open(test_pdf_file, "rb") as f:
            response = client.post(
                "/api/contracts/upload",
                files={"file": ("test_nda.pdf", f, "application/pdf")},
                data={"contract_name": "Test NDA Upload", "contract_type": "NDA"}
            )

        assert response.status_code == 201
        data = response.json()
        assert data["status"] == "success"
        assert data["contract_name"] == "Test NDA Upload"
        assert data["file_type"] == "PDF"
        assert data["page_count"] == 2
        assert data["chunk_count"] >= 1
    finally:
        app.dependency_overrides.clear()
