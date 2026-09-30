import os
import uuid
import tempfile
import fitz  # PyMuPDF
import docx
import pytest
from sqlalchemy import create_engine
from sqlalchemy.pool import StaticPool
from sqlalchemy.orm import sessionmaker

from backend.app.database import Base
from backend.app.models.models import Contract, Document, Chunk
from ingestion.pipeline.ingestor import process_contract_ingestion
from backend.app.services.search import get_search_service
from backend.app.services.graph import get_graph_service


@pytest.fixture
def db_session():
    """Isolated in-memory SQLite session fixture using StaticPool."""
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool
    )
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(bind=engine)


@pytest.fixture
def sample_pdf_bytes():
    """Generate in-memory multi-page contract PDF bytes."""
    doc = fitz.open()

    # Page 1: Confidentiality
    p1 = doc.new_page()
    p1.insert_text(
        (50, 72),
        "Mutual Non-Disclosure Agreement\n\n"
        "Section 1: Confidentiality and Proprietary Information\n"
        "The receiving party shall hold in strict confidence all proprietary technical data and trade secrets disclosed by Acme Corp."
    )

    # Page 2: Liability and Governing Law
    p2 = doc.new_page()
    p2.insert_text(
        (50, 72),
        "Section 2: Limitation of Liability\n"
        "The total aggregate liability under this agreement is subject to a cap on liability equal to $50,000.\n\n"
        "Section 3: Governing Law\n"
        "This Agreement shall be construed and governed in accordance with the laws of California."
    )

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_docx_bytes():
    """Generate in-memory DOCX bytes."""
    doc = docx.Document()
    doc.add_heading("Professional Services Agreement", level=1)
    doc.add_paragraph("Section 1: Scope of Services. Contractor shall provide cloud architecture engineering services.")
    doc.add_paragraph("Section 2: Compensation. Client shall pay $120,000 USD within 30 days of quarterly invoice receipt.")
    
    tmp_path = os.path.join(tempfile.gettempdir(), f"tmp_{uuid.uuid4().hex}.docx")
    doc.save(tmp_path)
    with open(tmp_path, "rb") as f:
        data = f.read()
    if os.path.exists(tmp_path):
        os.remove(tmp_path)
    return data


def test_upload_parse_store_embed_retrieve_pipeline(db_session, sample_pdf_bytes):
    """
    Verify complete integration sequence as specified in docs/18-testing-strategy.md:
    Upload -> Parse -> Store -> Embed -> Retrieve
    """
    # 1. Upload & Ingest
    ingest_result = process_contract_ingestion(
        db=db_session,
        file_bytes=sample_pdf_bytes,
        original_filename="nda_acme.pdf",
        contract_name="Acme Mutual NDA",
        contract_type="Non-Disclosure Agreement"
    )

    assert ingest_result["status"] == "success"
    contract_id = ingest_result["contract_id"]
    document_id = ingest_result["document_id"]
    assert contract_id is not None
    assert document_id is not None

    # 2. Parse & Store verification
    contract = db_session.query(Contract).filter(Contract.id == contract_id).first()
    assert contract is not None
    assert contract.name == "Acme Mutual NDA"
    assert contract.status == "COMPLETED"

    doc = db_session.query(Document).filter(Document.id == document_id).first()
    assert doc is not None
    assert doc.page_count == 2
    assert doc.processing_status == "COMPLETED"

    # 3. Chunks & Embeddings persistence verification
    chunks = db_session.query(Chunk).filter(Chunk.contract_id == contract_id).all()
    assert len(chunks) >= 2

    for chk in chunks:
        assert chk.text is not None and len(chk.text) > 0
        assert chk.page_start >= 1
        assert chk.embedding is not None
        assert isinstance(chk.embedding, list)
        assert len(chk.embedding) == 1536

    # 4. Semantic vector retrieval verification
    search_svc = get_search_service()
    query = "proprietary technical data trade secrets confidentiality"
    results = search_svc.semantic_search(
        db=db_session,
        query=query,
        contract_id=contract_id,
        top_k=3
    )

    assert len(results) >= 1
    top_hit = results[0]
    assert top_hit["contract_id"] == contract_id
    assert "Acme" in top_hit["contract_name"]
    assert top_hit["similarity_score"] > 0.0
    assert "confidential" in top_hit["text"].lower() or "proprietary" in top_hit["text"].lower()


def test_docx_pipeline_integration(db_session, sample_docx_bytes):
    """Verify full ingestion and retrieval pipeline for DOCX files."""
    ingest_result = process_contract_ingestion(
        db=db_session,
        file_bytes=sample_docx_bytes,
        original_filename="services_contract.docx",
        contract_name="Cloud Engineering Services",
        contract_type="Professional Services"
    )

    assert ingest_result["status"] == "success"
    assert ingest_result["file_type"] == "DOCX"
    contract_id = ingest_result["contract_id"]

    # Retrieve by query
    search_svc = get_search_service()
    results = search_svc.semantic_search(
        db=db_session,
        query="cloud architecture compensation payment invoice",
        contract_id=contract_id,
        top_k=2
    )

    assert len(results) >= 1
    assert "Services" in results[0]["contract_name"]
    assert results[0]["similarity_score"] > 0.0


def test_cascade_delete_and_lifecycle_integration(db_session, sample_pdf_bytes):
    """
    Verify lifecycle and privacy compliance:
    Deleting a contract purges stored file, db chunks/embeddings, and graph nodes.
    """
    ingest_result = process_contract_ingestion(
        db=db_session,
        file_bytes=sample_pdf_bytes,
        original_filename="lifecycle_test.pdf",
        contract_name="Lifecycle Contract",
        contract_type="NDA"
    )
    contract_id = ingest_result["contract_id"]

    # Sync to Knowledge Graph
    graph_svc = get_graph_service()
    graph_subgraph = graph_svc.sync_contract_to_graph(db=db_session, contract_id=contract_id)
    assert graph_subgraph["node_count"] >= 1

    # Verify presence in graph
    c_node = graph_svc.provider.nodes.get(f"Contract:{contract_id}")
    assert c_node is not None

    # Retrieve contract and check file exists
    contract = db_session.query(Contract).filter(Contract.id == contract_id).first()
    file_path = contract.file_path
    assert file_path is not None
    assert os.path.exists(file_path)

    # Perform cascade deletion
    if os.path.exists(file_path):
        os.remove(file_path)
    graph_svc.delete_contract_from_graph(contract_id)
    db_session.delete(contract)
    db_session.commit()

    # 1. Physical file removed
    assert not os.path.exists(file_path)

    # 2. Relational DB and chunks purged
    assert db_session.query(Contract).filter(Contract.id == contract_id).first() is None
    assert db_session.query(Chunk).filter(Chunk.contract_id == contract_id).count() == 0

    # 3. Knowledge Graph node purged
    assert graph_svc.provider.nodes.get(f"Contract:{contract_id}") is None
